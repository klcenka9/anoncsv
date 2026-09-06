import csv
import json
from pathlib import Path

from anoncsv import (
    anonymize_rows,
    anonymize_text,
    fake_email,
    fake_ip,
    fake_phone,
    parse_mapping,
    process_json_file,
    process_file,
    process_xlsx_file,
)


def test_email_is_deterministic_and_not_original():
    first = fake_email("anna@example.com", "test-salt")
    assert first == fake_email("anna@example.com", "test-salt")
    assert first != "anna@example.com"
    assert first.endswith(".invalid")
    assert "_" not in first.split("@", 1)[1]


def test_phone_keeps_shape_and_changes_value():
    result = fake_phone("+420 777 123 456", "test-salt")
    assert result.startswith("+")
    assert len(result) == 13
    assert result != "+420 777 123 456"


def test_ip_uses_documentation_range():
    assert fake_ip("192.168.1.20", "test-salt").startswith("198.51.100.")


def test_anonymize_rows_preserves_empty_values():
    rows = [{"email": "a@example.com", "note": "", "id": "42"}]
    result = anonymize_rows(rows, {"email": "email", "id": "token"}, "salt")
    assert result[0]["note"] == ""
    assert result[0]["email"] != rows[0]["email"]
    assert result[0]["id"] != "42"


def test_process_file_infers_columns_and_preserves_header(tmp_path: Path):
    source = tmp_path / "input.csv"
    target = tmp_path / "output.csv"
    source.write_text(
        "name,email,city\nAnna Nováková,anna@example.com,Praha\n",
        encoding="utf-8",
    )
    count, fields = process_file(source, target, {}, "salt", ",")
    with target.open(encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert count == 1
    assert fields == {"name": "name", "email": "email"}
    assert rows[0]["city"] == "Praha"
    assert rows[0]["email"] != "anna@example.com"


def test_parse_mapping_rejects_unknown_type():
    try:
        parse_mapping(["email=unknown"])
    except ValueError as error:
        assert "Typ musí být" in str(error)
    else:
        raise AssertionError("Expected ValueError")


def test_anonymize_text_returns_csv_for_web_ui():
    output, count, fields = anonymize_text(
        "email,city\na@example.com,Praha\n", {}, "salt"
    )
    assert count == 1
    assert fields == {"email": "email"}
    assert "Praha" in output
    assert "a@example.com" not in output


def test_web_page_contains_local_file_picker():
    from webapp import render_page

    page = render_page().decode("utf-8")
    assert 'type="file"' in page
    assert "file.text()" in page


def test_web_page_can_render_anonymized_preview():
    from webapp import render_page

    page = render_page(result='<section class="result">preview</section>').decode("utf-8")
    assert 'class="result"' in page
    assert "preview" in page


def test_anonymize_text_rejects_invalid_delimiter():
    try:
        anonymize_text("email;city\na@example.com;Praha\n", {}, "salt", ";;")
    except ValueError as error:
        assert "jeden znak" in str(error)
    else:
        raise AssertionError("Expected ValueError")


def test_process_json_file_anonymizes_array(tmp_path: Path):
    source = tmp_path / "input.json"
    target = tmp_path / "output.json"
    source.write_text(
        json.dumps([{"email": "a@example.com", "city": "Praha"}]),
        encoding="utf-8",
    )
    count, fields = process_json_file(source, target, {}, "salt")
    result = json.loads(target.read_text(encoding="utf-8"))
    assert count == 1
    assert fields == {"email": "email"}
    assert result[0]["city"] == "Praha"
    assert result[0]["email"] != "a@example.com"


def test_process_xlsx_file_anonymizes_first_sheet(tmp_path: Path):
    from openpyxl import Workbook, load_workbook

    source = tmp_path / "input.xlsx"
    target = tmp_path / "output.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["email", "city"])
    sheet.append(["a@example.com", "Praha"])
    workbook.save(source)
    count, fields = process_xlsx_file(source, target, {}, "salt")
    result = load_workbook(target).active
    assert count == 1
    assert fields == {"email": "email"}
    assert result.cell(2, 2).value == "Praha"
    assert result.cell(2, 1).value != "a@example.com"
